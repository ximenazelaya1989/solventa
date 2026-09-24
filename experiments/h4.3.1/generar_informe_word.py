from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "output" / "docx"
OUT = OUT_DIR / "Informe_Completo_Experimento_H4_3_1.docx"
CHART_DIR = OUT_DIR / "assets"

BLUE = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


vertical = [
    ["1 / 1 GiB", 200, 199.88, 11.11, 57.65, 0.000, 0],
    ["1 / 1 GiB", 500, 493.56, 634.84, 1571.83, 0.000, 95],
    ["1 / 1 GiB", 800, 799.85, 29.56, 62.37, 0.000, 0],
    ["1 / 1 GiB", 1100, 919.85, 1637.24, 1715.96, 0.000, 1852],
    ["1 / 1 GiB", 1400, 877.55, 4187.90, 4904.78, 0.000, 6993],
    ["2 / 2 GiB", 200, 199.99, 6.32, 14.46, 0.000, 0],
    ["2 / 2 GiB", 500, 499.91, 6.72, 49.47, 0.000, 0],
    ["2 / 2 GiB", 800, 799.84, 21.20, 40.71, 0.000, 0],
    ["2 / 2 GiB", 1100, 1001.19, 1019.29, 1431.81, 0.000, 1096],
    ["2 / 2 GiB", 1400, 965.38, 3064.70, 3236.13, 0.000, 5675],
]

horizontal = [
    [3, "1 / 1 GiB", 200, 199.96, 9.54, 15.81, 0.000, 0],
    [3, "1 / 1 GiB", 500, 499.86, 8.48, 19.30, 0.000, 0],
    [3, "1 / 1 GiB", 800, 799.59, 9.49, 30.08, 0.000, 0],
    [3, "1 / 1 GiB", 1100, 1099.59, 8.25, 18.40, 0.000, 0],
    [3, "1 / 1 GiB", 1400, 1399.42, 9.56, 40.65, 0.000, 0],
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(value)
        set_cell_shading(cell, BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell)
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(font_size)
    for ridx, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            if ridx % 2:
                set_cell_shading(cells[i], PALE_BLUE)
            if widths:
                cells[i].width = Inches(widths[i])
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def fmt(value, decimals=2):
    if isinstance(value, int):
        return f"{value:,}".replace(",", ".")
    return f"{value:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def create_charts():
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    rates = [200, 500, 800, 1100, 1400]
    v1 = [r[2] for r in vertical[:5]]
    v2 = [r[2] for r in vertical[5:]]
    h = [r[3] for r in horizontal]
    throughput = CHART_DIR / "throughput.png"
    draw_line_chart(
        throughput,
        "Throughput logrado por configuración",
        "TPS ofrecidos",
        "TPS logrados",
        rates,
        [("Carga ofrecida", rates, "#777777"), ("Vertical 1 CPU 1 GiB", v1, "#C0504D"),
         ("Vertical 2 CPU 2 GiB", v2, "#F79646"), ("Horizontal 3 instancias", h, "#1F4E79")],
        1500,
    )
    p95 = [r[4] for r in horizontal]
    p99 = [r[5] for r in horizontal]
    latency = CHART_DIR / "latencias_horizontal.png"
    draw_line_chart(
        latency,
        "Latencias horizontales frente a los límites",
        "TPS ofrecidos",
        "Latencia ms",
        rates,
        [("p95", p95, "#1F4E79"), ("p99", p99, "#70AD47"),
         ("Límite p95 400 ms", [400] * 5, "#C0504D"), ("Límite p99 800 ms", [800] * 5, "#9C0006")],
        850,
    )
    return throughput, latency


def draw_line_chart(path, title, xlabel, ylabel, xs, series, ymax):
    width, height = 1600, 820
    margin_left, margin_right, margin_top, margin_bottom = 150, 80, 95, 150
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(font_path, 28)
    small = ImageFont.truetype(font_path, 23)
    bold = ImageFont.truetype(bold_path, 34)
    plot_w = width - margin_left - margin_right
    plot_h = height - margin_top - margin_bottom
    x_min, x_max = min(xs), max(xs)

    def xp(v): return margin_left + (v - x_min) / (x_max - x_min) * plot_w
    def yp(v): return margin_top + plot_h - v / ymax * plot_h

    draw.text((width / 2, 26), title, fill="#000000", font=bold, anchor="ma")
    for i in range(6):
        value = ymax * i / 5
        y = yp(value)
        draw.line((margin_left, y, width - margin_right, y), fill="#D9D9D9", width=2)
        draw.text((margin_left - 18, y), f"{value:.0f}", fill="#444444", font=small, anchor="rm")
    draw.line((margin_left, margin_top, margin_left, margin_top + plot_h), fill="#333333", width=3)
    draw.line((margin_left, margin_top + plot_h, width - margin_right, margin_top + plot_h), fill="#333333", width=3)
    for x in xs:
        px = xp(x)
        draw.text((px, margin_top + plot_h + 16), f"{x}", fill="#333333", font=small, anchor="ma")
    draw.text((margin_left + plot_w / 2, height - 22), xlabel, fill="#222222", font=font, anchor="ms")
    label_box = Image.new("RGBA", (340, 60), (255, 255, 255, 0))
    label_draw = ImageDraw.Draw(label_box)
    label_draw.text((170, 30), ylabel, fill="#222222", font=font, anchor="mm")
    label_box = label_box.rotate(90, expand=True)
    image.paste(label_box, (20, int(margin_top + plot_h / 2 - label_box.height / 2)), label_box)
    for label, values, color in series:
        points = [(xp(x), yp(y)) for x, y in zip(xs, values)]
        draw.line(points, fill=color, width=6)
        for point in points:
            draw.ellipse((point[0] - 7, point[1] - 7, point[0] + 7, point[1] + 7), fill=color)
    legend_y = height - 83
    legend_x = margin_left
    for label, _values, color in series:
        draw.line((legend_x, legend_y, legend_x + 42, legend_y), fill=color, width=6)
        draw.text((legend_x + 53, legend_y), label, fill="#222222", font=small, anchor="lm")
        legend_x += int(draw.textlength(label, font=small)) + 105
    image.save(path, dpi=(180, 180))


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    throughput_chart, latency_chart = create_charts()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (("Title", 26, 0, 10), ("Heading 1", 17, 14, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11, 8, 4)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Title" or True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    old_border = title_ppr.find(qn("w:pBdr"))
    if old_border is not None:
        title_ppr.remove(old_border)

    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.text = "Solventa  Experimento H4.3.1  Página "
        footer.runs[0].font.size = Pt(9)
        add_page_number(footer)

    # Portada
    doc.add_paragraph("UNIVERSIDAD DE LOS ANDES").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("ISIS2212")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    doc.add_paragraph("Sprint 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("\n\n")
    title = doc.add_paragraph("Informe del experimento individual de escalabilidad H4.3.1", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Reprocesamiento masivo de perfiles con fuentes Open Data")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(15)
    sub.runs[0].bold = True
    doc.add_paragraph("\n")
    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("Estudiante", "Beraly Ventura"),
        ("Repositorio", "github.com/ximenazelaya1989/solventa"),
        ("Rama", "experimento-escalabilidad-h4-3-1"),
        ("Commit del informe", "7a7d0f4"),
        ("Fecha", "24 de septiembre de 2026"),
    ]
    for i, (k, v) in enumerate(values):
        meta.cell(i, 0).text = k
        meta.cell(i, 1).text = v
        for j in range(2):
            set_cell_border(meta.cell(i, j))
            set_cell_margins(meta.cell(i, j), 120, 120, 120, 120)
            meta.cell(i, j).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(meta.cell(i, 0), LIGHT_BLUE)
        meta.cell(i, 0).paragraphs[0].runs[0].bold = True
    doc.add_paragraph("\n")
    conclusion = doc.add_paragraph()
    conclusion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = conclusion.add_run("Conclusión principal")
    run.bold = True
    conclusion.add_run("\nEl escalamiento horizontal cumplió los criterios hasta 1.400 TPS; una sola instancia no los sostuvo en la carga máxima.")
    doc.add_page_break()

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Este informe presenta la implementación y ejecución en AWS del experimento individual H4.3.1. "
        "El objetivo fue evaluar si el servicio de reprocesamiento masivo de perfiles puede acercarse a la tasa necesaria para procesar 10 millones de perfiles en dos horas, comparando escalamiento vertical y horizontal. "
        "La carga se incrementó de 200 a 1.400 transacciones por segundo y se midieron throughput, latencias p95 y p99, errores e iteraciones descartadas."
    )
    doc.add_paragraph(
        "La configuración horizontal de tres instancias de 1 CPU y 1 GiB detrás de un Application Load Balancer sostuvo 1.399,42 TPS a la carga máxima, con p95 de 9,56 ms, p99 de 40,65 ms, 0 % de errores y cero descartes. "
        "Cumplió los criterios p95 menor o igual a 400 ms, p99 menor o igual a 800 ms y errores menores al 1 %. "
        "Las configuraciones verticales dejaron de cumplir a cargas altas. Por tanto, se recomienda escalamiento horizontal como táctica principal."
    )
    doc.add_heading("Contenido del entregable", level=2)
    for text in (
        "Código del flujo de reprocesamiento y endpoint de salud.",
        "Escenarios de carga reproducibles con k6 y configuraciones Docker.",
        "Ejecución vertical y horizontal en AWS Academy.",
        "Resultados, análisis, decisión y limitaciones.",
        "Registro de evidencias y guía para el video de demostración.",
    ):
        add_bullet(doc, text)

    doc.add_heading("1 Contexto y atributo de calidad", level=1)
    doc.add_heading("1.1 ASR H4.3.1", level=2)
    doc.add_paragraph(
        "Ante la incorporación de nuevas fuentes de Open Data, el sistema debe reprocesar masivamente los perfiles de riesgo. "
        "La referencia del experimento es procesar al menos 10 millones de perfiles en menos de dos horas. Esto exige una tasa promedio de 1.388,89 perfiles por segundo; por ello, la carga máxima se redondeó a 1.400 TPS."
    )
    doc.add_heading("1.2 Hipótesis", level=2)
    doc.add_paragraph(
        "El escalamiento horizontal mediante múltiples instancias detrás de un balanceador mantendrá un throughput útil mayor y latencias más estables que una sola instancia escalada verticalmente cuando la carga alcance 1.400 TPS."
    )
    doc.add_heading("1.3 Criterios de aceptación", level=2)
    add_table(doc, ["Métrica", "Criterio", "Interpretación"], [
        ["Throughput", "Cercano a la carga ofrecida", "Debe sostener hasta 1.400 TPS sin descartes"],
        ["Latencia p95", "≤ 400 ms", "El 95 % de las solicitudes no supera el límite"],
        ["Latencia p99", "≤ 800 ms", "El 99 % de las solicitudes no supera el límite"],
        ["Errores", "< 1 %", "Solicitudes HTTP fallidas"],
        ["Descartadas", "0 deseable", "k6 pudo iniciar toda la carga programada"],
    ], [1.15, 1.55, 4.0], 9)

    doc.add_heading("2 Correspondencia con la aplicación", level=1)
    doc.add_paragraph(
        "El repositorio original contenía la entidad PerfilRiesgo y una operación de lectura de score, pero no incluía el flujo completo de perfilamiento ni un canal en línea listo para medir. "
        "Para mantener el experimento simple y alineado con la implementación real, se agregó POST /perfilamiento/reprocesar. El endpoint recibe señales Open Data, calcula un score y ejecuta un upsert por clienteId sobre PostgreSQL. Una respuesta exitosa representa un perfil reprocesado."
    )
    doc.add_paragraph(
        "El experimento valida la táctica arquitectónica y el arnés de carga. No demuestra la capacidad definitiva de una futura lógica de negocio más compleja ni la ausencia de degradación de un canal en línea que todavía no existe en el repositorio. Esta limitación se conserva explícitamente para evitar conclusiones superiores a la evidencia."
    )
    doc.add_heading("2.1 Implementación incluida", level=2)
    add_table(doc, ["Componente", "Propósito"], [
        ["POST /perfilamiento/reprocesar", "Operación idempotente de reprocesamiento sobre PostgreSQL"],
        ["GET /health", "Verificación de salud utilizada por el balanceador"],
        ["k6/reprocesamiento.js", "Carga de tasa constante y recolección de métricas"],
        ["Docker Compose vertical", "Una instancia con CPU y memoria configurables"],
        ["Docker Compose horizontal", "Tres instancias y balanceo Nginx para el ensayo local"],
        ["run.sh", "Orquestación de niveles y archivos de resumen"],
    ], [2.3, 4.4], 9)

    doc.add_heading("3 Metodología", level=1)
    doc.add_heading("3.1 Diseño experimental", level=2)
    doc.add_paragraph(
        "Se aplicaron cinco niveles de carga en orden ascendente: 200, 500, 800, 1.100 y 1.400 TPS. Cada nivel se ejecutó durante 15 segundos desde AWS CloudShell con k6 0.55.0. "
        "La duración se redujo para proteger el presupuesto de AWS Academy; permite comparar las tácticas, pero no reemplaza una prueba de estabilidad de dos horas."
    )
    doc.add_heading("3.2 Variables", level=2)
    add_table(doc, ["Tipo", "Variable", "Valores o medición"], [
        ["Independiente", "Táctica", "Vertical y horizontal"],
        ["Independiente", "Carga", "200, 500, 800, 1.100 y 1.400 TPS"],
        ["Independiente", "Recursos", "1 CPU/1 GiB; 2 CPU/2 GiB; 3 × 1 CPU/1 GiB"],
        ["Dependiente", "Rendimiento", "TPS logrados, p95 y p99"],
        ["Dependiente", "Confiabilidad de la carga", "Errores e iteraciones descartadas"],
        ["Controlada", "Versión y datos", "Mismo commit, endpoint, región y PostgreSQL"],
    ], [1.25, 1.8, 3.65], 9)
    doc.add_heading("3.3 Procedimiento", level=2)
    for text in (
        "Desplegar PostgreSQL 16 y la misma versión de la aplicación.",
        "Verificar GET /health antes de iniciar las mediciones.",
        "Ejecutar la configuración vertical con una instancia y aumentar CPU y memoria.",
        "Ejecutar la configuración horizontal con tres instancias saludables detrás del ALB.",
        "Aplicar los cinco niveles de carga desde CloudShell y conservar cada resumen JSON.",
        "Consolidar métricas, verificar criterios y comparar throughput útil.",
    ):
        add_bullet(doc, text)

    doc.add_heading("4 Infraestructura de ejecución", level=1)
    add_table(doc, ["Elemento", "Configuración observada"], [
        ["Cuenta", "AWS Academy 289908319652"],
        ["Región", "us-east-1"],
        ["Aplicación", "Tres EC2 t2.medium distribuidas en zonas de disponibilidad"],
        ["Base de datos", "PostgreSQL 16 en una cuarta EC2"],
        ["Balanceador", "Application Load Balancer Cheapest-alb"],
        ["Grupo de destino", "Cheapest-tg con tres destinos saludables"],
        ["Health check", "GET /health"],
        ["Generador", "k6 0.55.0 desde AWS CloudShell"],
        ["Duración", "15 segundos por nivel"],
    ], [2.0, 4.7], 9)
    doc.add_paragraph(
        "Las instancias se detuvieron después de obtener los resultados para reducir el consumo de créditos. Los resultados y el informe permanecen versionados en GitHub; detener la infraestructura no modifica las mediciones ya obtenidas."
    )

    doc.add_page_break()
    doc.add_heading("5 Resultados", level=1)
    doc.add_heading("5.1 Escalamiento vertical", level=2)
    v_rows = [[r[0], fmt(r[1], 0), fmt(r[2]), fmt(r[3]), fmt(r[4]), fmt(r[5], 3), fmt(r[6], 0)] for r in vertical]
    add_table(doc, ["CPU memoria", "TPS ofrecidos", "TPS logrados", "p95 ms", "p99 ms", "Errores %", "Descartadas"], v_rows, [1.0, .9, .9, .85, .85, .8, .85], 7.3)
    doc.add_paragraph(
        "Con 2 CPU y 2 GiB la instancia se mantuvo dentro de los criterios hasta 800 TPS. A 1.100 y 1.400 TPS excedió p95 y p99, perdió throughput e introdujo iteraciones descartadas. "
        "En 1 CPU y 1 GiB se observó un valor anómalo a 500 TPS y recuperación a 800 TPS; se conserva tal como fue medido y no se corrige artificialmente."
    )
    doc.add_heading("5.2 Escalamiento horizontal", level=2)
    h_rows = [[fmt(r[0], 0), r[1], fmt(r[2], 0), fmt(r[3]), fmt(r[4]), fmt(r[5]), fmt(r[6], 3), fmt(r[7], 0)] for r in horizontal]
    add_table(doc, ["Instancias", "CPU memoria c/u", "TPS ofrecidos", "TPS logrados", "p95 ms", "p99 ms", "Errores %", "Descartadas"], h_rows, [.7, 1.05, .8, .8, .65, .65, .65, .72], 6.8)
    doc.add_paragraph(
        "La configuración horizontal cumplió todos los criterios en los cinco niveles. A 1.400 TPS logró 1.399,42 TPS, p95 de 9,56 ms, p99 de 40,65 ms, 0 % de errores y ninguna iteración descartada."
    )

    doc.add_heading("6 Análisis comparativo", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(throughput_chart), width=Inches(6.35))
    cap = doc.add_paragraph("Figura 1  Throughput logrado por configuración")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_paragraph(
        "La separación aparece en 1.100 TPS y se amplía a 1.400 TPS. La alternativa horizontal permanece cerca de la carga ofrecida; las instancias verticales se saturan y dejan de iniciar parte de las iteraciones."
    )
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(latency_chart), width=Inches(6.35))
    cap = doc.add_paragraph("Figura 2  Latencias horizontales frente a los límites")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_paragraph(
        "Las latencias horizontales se mantuvieron muy por debajo de ambos umbrales. El resultado respalda la hipótesis para el endpoint y el entorno probados. La tasa máxima equivale aproximadamente a 10,08 millones de perfiles en dos horas si puede sostenerse durante todo ese periodo."
    )
    doc.add_heading("6.1 Matriz de aceptación a 1.400 TPS", level=2)
    add_table(doc, ["Configuración", "TPS logrados", "p95", "p99", "Errores", "Descartes", "Decisión"], [
        ["Vertical 1 CPU 1 GiB", "877,55", "4.187,90", "4.904,78", "0 %", "6.993", "No cumple"],
        ["Vertical 2 CPU 2 GiB", "965,38", "3.064,70", "3.236,13", "0 %", "5.675", "No cumple"],
        ["Horizontal 3 instancias", "1.399,42", "9,56", "40,65", "0 %", "0", "Cumple"],
    ], [1.55, .85, .75, .75, .65, .7, .8], 7.8)

    doc.add_heading("7 Decisión arquitectónica", level=1)
    doc.add_paragraph(
        "La hipótesis queda respaldada: para este flujo de reprocesamiento, tres instancias pequeñas detrás del balanceador superan claramente a una sola instancia aun después de duplicar sus recursos. "
        "Se recomienda utilizar escalamiento horizontal como táctica principal, conservar el escalamiento vertical como ajuste inicial y monitorear la base de datos para evitar que se convierta en el siguiente cuello de botella."
    )
    doc.add_heading("7.1 Riesgos y limitaciones", level=2)
    add_table(doc, ["Limitación", "Efecto", "Trabajo posterior"], [
        ["Duración de 15 segundos", "No demuestra estabilidad durante dos horas", "Ejecutar una prueba prolongada con presupuesto controlado"],
        ["PostgreSQL en EC2", "No representa exactamente una implementación con RDS", "Repetir con RDS y métricas de CPU, conexiones e IOPS"],
        ["Lógica mínima de negocio", "La lógica futura puede modificar latencias", "Repetir al completar el perfilamiento definitivo"],
        ["Sin canal en línea", "No se pudo medir degradación paralela", "Agregar prueba concurrente cuando exista ese flujo"],
        ["Una corrida por nivel en AWS", "No cuantifica variabilidad", "Realizar al menos tres repeticiones por configuración"],
    ], [1.55, 2.25, 2.9], 8.4)

    doc.add_heading("8 Evidencias y trazabilidad", level=1)
    add_table(doc, ["Evidencia", "Ubicación", "Estado"], [
        ["Código implementado", "Rama experimento-escalabilidad-h4-3-1", "Disponible en GitHub"],
        ["Commit desplegado en AWS", "12bc015", "Versionado"],
        ["Commit del informe final", "7a7d0f4", "Versionado"],
        ["Resultados exactos", "experiments/h4.3.1/AWS_RESULTADOS.tsv", "Versionados"],
        ["Análisis detallado", "experiments/h4.3.1/RESULTADOS.md", "Versionado"],
        ["JSON crudos", "CloudShell ~/h431-aws-results.tgz", "Conservados en la sesión AWS"],
        ["Infraestructura observada", "ALB con tres destinos saludables", "Verificada durante la ejecución"],
        ["Video", "Enlace por agregar después de grabarlo", "Pendiente de la autora"],
    ], [1.55, 3.55, 1.6], 8.5)
    doc.add_heading("8.1 Enlaces de entrega", level=2)
    doc.add_paragraph("Repositorio: https://github.com/ximenazelaya1989/solventa")
    doc.add_paragraph("Rama: https://github.com/ximenazelaya1989/solventa/tree/experimento-escalabilidad-h4-3-1")
    doc.add_paragraph("Pull Request: opcional; crear únicamente para revisión y no fusionar antes de la aprobación del equipo.")

    doc.add_heading("9 Reproducción", level=1)
    doc.add_paragraph("Desde la raíz del repositorio, validar el código con:")
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    r = p.add_run("npm ci\nnpm test\nnpm run build")
    r.font.name = "Courier New"
    r.font.size = Pt(9)
    doc.add_paragraph("Ensayo vertical local:")
    p = doc.add_paragraph()
    r = p.add_run("API_CPUS=2 API_MEMORY=2g npm run experiment:h431 -- vertical")
    r.font.name = "Courier New"; r.font.size = Pt(9)
    doc.add_paragraph("Ensayo horizontal local:")
    p = doc.add_paragraph()
    r = p.add_run("INSTANCE_CPUS=1 INSTANCE_MEMORY=1g npm run experiment:h431 -- horizontal")
    r.font.name = "Courier New"; r.font.size = Pt(9)
    doc.add_paragraph("Ejecución de un nivel en AWS:")
    p = doc.add_paragraph()
    r = p.add_run('BASE_URL="http://DNS-DEL-ALB" RATE=1400 DURATION=15s SUMMARY_FILE=summary-1400-tps.json k6 run experiments/h4.3.1/k6/reprocesamiento.js')
    r.font.name = "Courier New"; r.font.size = Pt(8.5)
    doc.add_paragraph(
        "Para comparaciones válidas se debe conservar la misma versión, región, base de datos y generador; limpiar o restaurar los datos entre configuraciones y permitir un periodo de enfriamiento."
    )

    doc.add_heading("10 Guía del video de máximo cinco minutos", level=1)
    add_table(doc, ["Tiempo", "Contenido"], [
        ["0:00 a 0:30", "Presentar H4.3.1, la meta de 10 millones en dos horas y los límites p95/p99"],
        ["0:30 a 1:30", "Mostrar la rama, el endpoint, health check, k6 y configuraciones de escalamiento"],
        ["1:30 a 2:30", "Mostrar AWS: instancias, ALB y tres destinos saludables"],
        ["2:30 a 3:30", "Ejecutar o mostrar una ejecución real de k6 y explicar las métricas"],
        ["3:30 a 4:30", "Comparar las tablas vertical y horizontal"],
        ["4:30 a 5:00", "Concluir que horizontal cumple y explicar las limitaciones"],
    ], [1.2, 5.5], 9)
    doc.add_paragraph(
        "Antes de grabar, iniciar temporalmente las instancias si se necesita mostrar la infraestructura funcionando. Al finalizar, volver a detenerlas. No mostrar credenciales, claves, tokens ni información sensible de la cuenta."
    )

    doc.add_heading("11 Conclusiones", level=1)
    for text in (
        "El arnés experimental y el flujo mínimo de reprocesamiento quedaron implementados en una rama independiente; main no fue modificado.",
        "El escalamiento vertical mejoró la capacidad, pero no cumplió los percentiles ni sostuvo 1.400 TPS.",
        "El escalamiento horizontal sostuvo 1.399,42 TPS con p95 de 9,56 ms y p99 de 40,65 ms, sin errores ni descartes.",
        "La evidencia respalda el escalamiento horizontal como táctica principal para este flujo.",
        "La validación definitiva requiere una prueba prolongada, varias repeticiones y la lógica completa del negocio.",
    ):
        add_bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("Anexo A Archivos principales", level=1)
    add_table(doc, ["Archivo", "Contenido"], [
        ["experiments/h4.3.1/README.md", "Diseño, criterios y reproducción"],
        ["experiments/h4.3.1/RESULTADOS.md", "Resultados y análisis completos"],
        ["experiments/h4.3.1/AWS_RESULTADOS.tsv", "Valores fuente de las mediciones AWS"],
        ["experiments/h4.3.1/k6/reprocesamiento.js", "Escenario de carga"],
        ["output/pdf/Informe_Experimento_H4_3_1.pdf", "Informe ejecutivo PDF"],
        ["output/docx/Informe_Completo_Experimento_H4_3_1.docx", "Informe académico completo"],
    ], [3.4, 3.3], 9)

    props = doc.core_properties
    props.title = "Informe completo del experimento H4.3.1"
    props.subject = "Escalabilidad vertical y horizontal para reprocesamiento masivo de perfiles"
    props.author = "Beraly Ventura"
    props.keywords = "H4.3.1, escalabilidad, k6, AWS, p95, p99, throughput"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
